from pathlib import Path
import csv
import json
from typing import Any


class DatasetLoader:
    """
    Loads datasets from the DataGuard datasets directory.

    Supported formats:
    - CSV
    - JSON
    - XLSX

    Pandas is intentionally not used.
    """

    PROJECT_ROOT = Path(__file__).resolve().parents[4]
    DATASET_ROOT = PROJECT_ROOT / "datasets"

    @classmethod
    def load_dataset(
        cls,
        dataset_path: str,
    ) -> list[dict[str, Any]]:
        """
        Load a dataset and return it as a list of dictionaries.
        """

        full_path = (cls.DATASET_ROOT / dataset_path).resolve()
        dataset_root_resolved = cls.DATASET_ROOT.resolve()

        try:
            full_path.relative_to(dataset_root_resolved)
        except ValueError:
            raise PermissionError(
                f"Access denied: Path traversal detected for path '{dataset_path}'."
            )

        if not full_path.exists():
            raise FileNotFoundError(
                f"Dataset not found: {full_path}"
            )

        if not full_path.is_file():
            raise ValueError(
                f"Dataset path is not a file: {full_path}"
            )

        extension = full_path.suffix.lower()

        if extension == ".csv":
            return cls._load_csv(full_path)

        if extension == ".json":
            return cls._load_json(full_path)

        if extension == ".xlsx":
            return cls._load_excel(full_path)

        raise ValueError(
            f"Unsupported dataset format: {extension}"
        )

    @staticmethod
    def _load_csv(
        file_path: Path,
    ) -> list[dict[str, Any]]:
        """
        Load CSV using Python's built-in csv module.
        """

        with file_path.open(
            mode="r",
            encoding="utf-8-sig",
            newline="",
        ) as file:

            reader = csv.DictReader(file)

            if reader.fieldnames is None:
                raise ValueError(
                    "CSV file does not contain a header row."
                )

            return [dict(row) for row in reader]

    @staticmethod
    def _load_json(
        file_path: Path,
    ) -> list[dict[str, Any]]:
        """
        Load JSON dataset.

        Supported structures:

        [
            {"id": 1, "name": "Alice"}
        ]

        or:

        {
            "data": [
                {"id": 1, "name": "Alice"}
            ]
        }
        """

        with file_path.open(
            mode="r",
            encoding="utf-8",
        ) as file:

            data = json.load(file)

        if isinstance(data, list):

            if not all(
                isinstance(item, dict)
                for item in data
            ):
                raise ValueError(
                    "JSON dataset records must be objects."
                )

            return data

        if isinstance(data, dict):

            if "data" not in data:
                raise ValueError(
                    "JSON object must contain a 'data' field."
                )

            if not isinstance(data["data"], list):
                raise ValueError(
                    "JSON 'data' field must contain a list."
                )

            if not all(
                isinstance(item, dict)
                for item in data["data"]
            ):
                raise ValueError(
                    "JSON dataset records must be objects."
                )

            return data["data"]

        raise ValueError(
            "Unsupported JSON dataset structure."
        )

    @staticmethod
    def _load_excel(
        file_path: Path,
    ) -> list[dict[str, Any]]:
        """
        Load XLSX using openpyxl.
        """

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required to load XLSX datasets."
            ) from exc

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        workbook.close()

        if not rows:
            return []

        headers = rows[0]

        if not any(
            header is not None
            for header in headers
        ):
            raise ValueError(
                "Excel file does not contain a header row."
            )

        normalized_headers = [
            str(header).strip()
            if header is not None
            else f"column_{index + 1}"
            for index, header in enumerate(headers)
        ]

        result = []

        for row in rows[1:]:

            record = {}

            for index, column in enumerate(
                normalized_headers
            ):
                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                record[column] = value

            if any(
                value is not None
                for value in record.values()
            ):
                result.append(record)

        return result