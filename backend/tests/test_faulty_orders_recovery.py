"""
Regression Test Suite: Faulty Orders Recovery, Excel Export, and RBAC Enforcement
Validates the end-to-end recovery console against test/dataguard_faulty_orders.xlsx
per the enterprise DataGuard recovery specifications.
"""

import io
import os
import unittest
from uuid import uuid4
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.database.session import SessionLocal
from app.models.user import User
from app.core.security import create_access_token, hash_password
from app.services.inspection_finding_service import InspectionFindingService
from app.etl.extractor.dataset_loader import DatasetLoader
from app.agents.orchestrator import MultiAgentOrchestrator


class TestFaultyOrdersRecovery(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.db = SessionLocal()

        # Create test users: Admin and Data Engineer
        cls.admin_user = User(
            id=uuid4(),
            email=f"admin_recovery_{uuid4().hex[:6]}@dataguard.ai",
            full_name="Admin Recovery Officer",
            hashed_password=hash_password("AdminPass123!"),
            role="ADMIN",
            is_superuser=True,
            is_active=True,
        )
        cls.engineer_user = User(
            id=uuid4(),
            email=f"engineer_recovery_{uuid4().hex[:6]}@dataguard.ai",
            full_name="Staff Data Engineer",
            hashed_password=hash_password("EngPass123!"),
            role="DATA_ENGINEER",
            is_superuser=False,
            is_active=True,
        )
        cls.db.add(cls.admin_user)
        cls.db.add(cls.engineer_user)
        cls.db.commit()

        cls.admin_token = create_access_token({
            "sub": str(cls.admin_user.id),
            "role": "ADMIN",
            "email": cls.admin_user.email,
        })
        cls.engineer_token = create_access_token({
            "sub": str(cls.engineer_user.id),
            "role": "DATA_ENGINEER",
            "email": cls.engineer_user.email,
        })

        cls.dataset_path = "test/dataguard_faulty_orders.xlsx"
        cls.rows = DatasetLoader.load_dataset(cls.dataset_path)

        # Run multi-agent orchestrator inspection
        cls.orchestrator_result = MultiAgentOrchestrator().audit_dataset(
            rows=cls.rows,
            dataset_path=cls.dataset_path,
            generate_reports=False,
        )

        # Persist inspection run to database for audit linkage
        cls.inspection_run = InspectionFindingService.save_inspection(
            db=cls.db,
            dataset_path=cls.dataset_path,
            inspection_result=cls.orchestrator_result["inspection"],
            row_count=len(cls.rows),
            column_count=len(cls.rows[0]) if cls.rows else 0,
        )

    @classmethod
    def tearDownClass(cls):
        cls.db.close()

    def test_01_inspection_findings_and_proposals_integrity(self):
        """Verify that the 5 distinct defect types in dataguard_faulty_orders.xlsx are detected."""
        self.assertEqual(len(self.rows), 318, "Initial row count should be exactly 318")

        findings = self.orchestrator_result["inspection"]["findings"]
        self.assertEqual(len(findings), 5, f"Expected 5 findings, got {len(findings)}")

        finding_types = {(f.get("type"), f.get("column")) for f in findings}
        self.assertIn(("MISSING_VALUES", "customer_state"), finding_types)
        self.assertIn(("MISSING_VALUES", "delivery_days"), finding_types)
        self.assertIn(("MISSING_VALUES", "payment_type"), finding_types)
        self.assertIn(("DUPLICATE_RECORDS", None), finding_types)
        self.assertIn(("ML_ISOLATION_FOREST_ANOMALY", None), finding_types)

        candidates = self.orchestrator_result["recovery"]["candidates"]
        self.assertEqual(len(candidates), 5, f"Expected 5 candidate actions, got {len(candidates)}")

        # Verify role assignment on proposed actions
        role_map = {c["action_id"]: c.get("required_role") for c in candidates}
        self.assertEqual(role_map.get("ACT-001"), "DATA_ENGINEER")  # customer_state
        self.assertEqual(role_map.get("ACT-002"), "ADMIN")          # delivery_days (numeric)
        self.assertEqual(role_map.get("ACT-003"), "DATA_ENGINEER")  # payment_type
        self.assertEqual(role_map.get("ACT-004"), "DATA_ENGINEER")  # deduplication
        self.assertEqual(role_map.get("ACT-005"), "ADMIN")          # ML outlier quarantine

    def test_02_admin_full_recovery_and_excel_export(self):
        """ADMIN executes all 5 actions: 318 -> 292 rows, PASS verdict, clean XLSX/CSV/PDF."""
        payload = {
            "dataset_path": self.dataset_path,
            "audit_id": str(self.inspection_run.id),
            "actions": self.orchestrator_result["recovery"]["candidates"],
        }

        resp = self.client.post(
            "/agents/recovery/execute",
            json=payload,
            headers={"Authorization": f"Bearer {self.admin_token}"},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()

        self.assertEqual(data["status"], "RECOVERY_APPLIED")
        self.assertEqual(len(data["actions_executed"]), 5)
        self.assertEqual(len(data["actions_skipped"]), 0)

        verification = data["verification"]
        self.assertEqual(verification["verdict"], "PASS")
        self.assertEqual(verification["metrics"]["original_row_count"], 318)
        self.assertEqual(verification["metrics"]["remediated_row_count"], 292)
        self.assertEqual(verification["metrics"]["before_finding_count"], 5)
        self.assertEqual(verification["metrics"]["after_finding_count"], 0)
        self.assertEqual(verification["metrics"]["anomaly_reduction_rate_pct"], 100.0)

        run_id = data["recovery_run_id"]
        sha256 = data["remediated_file_hash"]

        # 1. Download Remediated Excel (.xlsx)
        resp_xlsx = self.client.get(
            f"/agents/recovery/download/{run_id}?format=xlsx",
            headers={"Authorization": f"Bearer {self.admin_token}"},
        )
        self.assertEqual(resp_xlsx.status_code, 200)
        self.assertIn("openxmlformats", resp_xlsx.headers["content-type"])

        # Parse Excel workbook with OpenPyXL
        wb = openpyxl.load_workbook(io.BytesIO(resp_xlsx.content))
        self.assertIn("Remediated Dataset", wb.sheetnames)
        self.assertIn("Recovery Audit Log", wb.sheetnames)
        self.assertIn("Remediation Actions", wb.sheetnames)

        ws_data = wb["Remediated Dataset"]
        self.assertEqual(ws_data.max_row, 293, "1 header row + 292 clean data rows")
        self.assertEqual(ws_data.max_column, 11)

        # Validate no cell contains "####" and numeric values are preserved
        for row in ws_data.iter_rows(values_only=True):
            for cell_val in row:
                if cell_val is not None:
                    self.assertNotIn("####", str(cell_val), "Found '####' overflow in Excel cell")

        # Validate column dimensions
        for col_letter, col_dim in ws_data.column_dimensions.items():
            if col_dim.width is not None:
                self.assertGreaterEqual(col_dim.width, 14, f"Column {col_letter} width too narrow")
                self.assertLessEqual(col_dim.width, 50, f"Column {col_letter} width exceeds limit")

        # 2. Download Remediated CSV (.csv)
        resp_csv = self.client.get(
            f"/agents/recovery/download/{run_id}?format=csv",
            headers={"Authorization": f"Bearer {self.admin_token}"},
        )
        self.assertEqual(resp_csv.status_code, 200)
        self.assertIn("text/csv", resp_csv.headers["content-type"])
        self.assertEqual(resp_csv.headers.get("x-file-sha256"), sha256)
        csv_lines = [line for line in resp_csv.text.strip().splitlines() if line.strip()]
        self.assertEqual(len(csv_lines), 293, "1 header + 292 data rows in CSV")

        # 3. Download Remediated PDF (.pdf)
        resp_pdf = self.client.get(
            f"/agents/recovery/download/{run_id}?format=pdf",
            headers={"Authorization": f"Bearer {self.admin_token}"},
        )
        self.assertEqual(resp_pdf.status_code, 200)
        self.assertIn("application/pdf", resp_pdf.headers["content-type"])
        self.assertTrue(resp_pdf.content.startswith(b"%PDF-"), "PDF missing standard magic header")

    def test_03_data_engineer_rbac_enforcement_and_partial_recovery(self):
        """DATA_ENGINEER executes permitted low-risk actions only: 318 -> 300 rows, FAIL verdict."""
        payload = {
            "dataset_path": self.dataset_path,
            "audit_id": str(self.inspection_run.id),
            "actions": self.orchestrator_result["recovery"]["candidates"],
        }

        resp = self.client.post(
            "/agents/recovery/execute",
            json=payload,
            headers={"Authorization": f"Bearer {self.engineer_token}"},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()

        self.assertEqual(data["status"], "RECOVERY_APPLIED")
        # 3 executed (ACT-001, ACT-003, ACT-004), 2 skipped (ACT-002, ACT-005)
        self.assertEqual(len(data["actions_executed"]), 3)
        self.assertEqual(len(data["actions_skipped"]), 2)

        skipped_ids = [s["action_id"] for s in data["actions_skipped"]]
        self.assertIn("ACT-002", skipped_ids)
        self.assertIn("ACT-005", skipped_ids)
        for s in data["actions_skipped"]:
            self.assertIn("ADMIN", s.get("reason", ""))

        verification = data["verification"]
        self.assertEqual(verification["verdict"], "FAIL")
        self.assertEqual(verification["metrics"]["original_row_count"], 318)
        self.assertEqual(verification["metrics"]["remediated_row_count"], 300)
        self.assertEqual(verification["metrics"]["before_finding_count"], 5)
        self.assertEqual(verification["metrics"]["after_finding_count"], 2)


if __name__ == "__main__":
    unittest.main()
